import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import re

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()
    
    # Page setup - Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # COVER PAGE
    # -------------------------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Woxsen University\n")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(160, 20, 30) # Woxsen Crimson
    
    run_sub = p.add_run("School of Technology\n")
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(16)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(70, 70, 70)

    run_dept = p.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\n\n")
    run_dept.font.name = 'Times New Roman'
    run_dept.font.size = Pt(13)
    run_dept.font.bold = True
    run_dept.font.color.rgb = RGBColor(90, 90, 90)

    p_pbl = doc.add_paragraph()
    p_pbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pbl = p_pbl.add_run("PROJECT BASED LEARNING REPORT\n")
    r_pbl.font.name = 'Times New Roman'
    r_pbl.font.size = Pt(15)
    r_pbl.font.bold = True
    r_pbl.font.underline = True

    r_course = p_pbl.add_run("Course: Database Management Systems\nAcademic Year: 2026–27\n\n")
    r_course.font.name = 'Times New Roman'
    r_course.font.size = Pt(13)
    r_course.font.italic = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(24)
    r_tlabel = p_title.add_run("PROJECT TITLE:\n")
    r_tlabel.font.size = Pt(13)
    r_tlabel.font.bold = True
    r_title = p_title.add_run("DESIGN AND IMPLEMENTATION OF A DATABASE MANAGEMENT SYSTEM FOR STUDENT ATTENDANCE AND LEAVE APPROVAL MANAGEMENT SYSTEM (AttendX)")
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(160, 20, 30)

    # Submission Box
    p_box = doc.add_paragraph()
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(20)
    
    table_cover = doc.add_table(rows=1, cols=2)
    table_cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_cover.autofit = False
    table_cover.columns[0].width = Inches(3.2)
    table_cover.columns[1].width = Inches(3.2)
    
    cell_l = table_cover.cell(0, 0)
    cell_r = table_cover.cell(0, 1)
    
    p_sub = cell_l.paragraphs[0]
    p_sub.add_run("Submitted by:\n").bold = True
    p_sub.add_run("Name: ").bold = True
    p_sub.add_run("Yasheswar Kataria\n")
    p_sub.add_run("Roll No: ").bold = True
    p_sub.add_run("25WU0101156\n")
    p_sub.add_run("Program: ").bold = True
    p_sub.add_run("B.Tech – CSE\n")
    p_sub.add_run("Section: ").bold = True
    p_sub.add_run("CSE - Panthers\n")
    
    p_gui = cell_r.paragraphs[0]
    p_gui.add_run("Under the Guidance of:\n").bold = True
    p_gui.add_run("Faculty Guide:\n").bold = True
    p_gui.add_run("Dr. Sahithi Chennamadhavuni\n")
    p_gui.add_run("Designation: ").bold = True
    p_gui.add_run("Associate Professor\n")
    p_gui.add_run("Dept. of Computer Science & Engg.\n")

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(40)
    r_date = p_date.add_run("October, 2026\nWoxsen University, Hyderabad")
    r_date.font.size = Pt(11)
    r_date.font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # CERTIFICATE & DECLARATION
    # -------------------------------------------------------------
    h_cert = doc.add_paragraph()
    h_cert.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h_cert.add_run("CERTIFICATE")
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.underline = True

    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_before = Pt(20)
    p_cert.paragraph_format.line_spacing = 1.3
    p_cert.add_run("This is to certify that the Project-Based Learning report entitled ")
    r_ctitle = p_cert.add_run("“DESIGN AND IMPLEMENTATION OF A DATABASE MANAGEMENT SYSTEM FOR STUDENT ATTENDANCE AND LEAVE APPROVAL MANAGEMENT SYSTEM (AttendX)”")
    r_ctitle.bold = True
    p_cert.add_run(" is a bonafide work carried out by ")
    p_cert.add_run("Mr. Yasheswar Kataria (Roll No: 25WU0101156)").bold = True
    p_cert.add_run(", student of B.Tech in Computer Science and Engineering at Woxsen University, during the Academic Year 2026–27, in partial fulfillment of the academic requirements for the course ")
    p_cert.add_run("Database Management Systems (DBMS)").bold = True
    p_cert.add_run(" under my supervision and guidance.")

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(80)
    p_sig.add_run("_________________________________________\n").bold = True
    p_sig.add_run("(Dr. Sahithi Chennamadhavuni)\n").bold = True
    p_sig.add_run("Associate Professor\nDepartment of Computer Science and Engineering\nSchool of Technology, Woxsen University")

    doc.add_page_break()

    # DECLARATION
    h_dec = doc.add_paragraph()
    h_dec.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h_dec.add_run("DECLARATION")
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.underline = True

    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_before = Pt(20)
    p_dec.paragraph_format.line_spacing = 1.3
    p_dec.add_run("I hereby declare that the project entitled ")
    p_dec.add_run("“DESIGN AND IMPLEMENTATION OF A DATABASE MANAGEMENT SYSTEM FOR STUDENT ATTENDANCE AND LEAVE APPROVAL MANAGEMENT SYSTEM (AttendX)”").bold = True
    p_dec.add_run(", submitted to the School of Technology, Woxsen University, Hyderabad, for the Project-Based Learning evaluation in the course Database Management Systems, is an authentic record of original work carried out by me under the guidance of ")
    p_dec.add_run("Dr. Sahithi Chennamadhavuni").bold = True
    p_dec.add_run(".\n\nI confirm that this work has not been previously submitted in part or full for the award of any other degree, diploma, or academic certificate in this or any other institution. All technical concepts, algorithms, literature references, and datasets used have been properly acknowledged and cited in accordance with university academic standards.\n\nI further declare that the content adheres strictly to academic integrity and ethical guidelines.")

    p_decsig = doc.add_paragraph()
    p_decsig.paragraph_format.space_before = Pt(60)
    p_decsig.add_run("Place: Hyderabad\nDate: 12th October 2026\n\n\n\n").bold = True
    p_decsig.add_run("_________________________________________\n").bold = True
    p_decsig.add_run("Signature of the Candidate\n").bold = True
    p_decsig.add_run("Name: Yasheswar Kataria\nRoll Number: 25WU0101156\nProgram: B.Tech – Computer Science and Engineering")

    doc.add_page_break()

    # -------------------------------------------------------------
    # PARSE AND INSERT REMAINING MD SECTIONS
    # -------------------------------------------------------------
    with open("docs/PROJECT_REPORT.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    # Split by chapters/sections
    # Let's clean the markdown and render headers, tables, code blocks and paragraphs nicely
    lines = md_text.split("\n")
    
    # We skip cover/cert/dec since we did them cleanly above
    # Find where Acknowledgement starts
    start_idx = 0
    for idx, l in enumerate(lines):
        if "## ACKNOWLEDGEMENT" in l:
            start_idx = idx
            break

    in_table = False
    table_lines = []
    in_code = False
    code_lines = []

    def flush_table(t_lines):
        if not t_lines: return
        rows = [row.strip().strip('|').split('|') for row in t_lines if '---' not in row]
        if not rows: return
        num_cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=num_cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Table Grid'
        
        for r_i, row in enumerate(rows):
            is_header = (r_i == 0)
            for c_i in range(num_cols):
                cell = t.cell(r_i, c_i)
                text = row[c_i].strip() if c_i < len(row) else ""
                cell.text = text
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                if is_header:
                    set_cell_background(cell, "E0E7FF") # Soft Indigo Header
                    for p_cell in cell.paragraphs:
                        for r_cell in p_cell.runs:
                            r_cell.font.bold = True
                            r_cell.font.size = Pt(10)
                            r_cell.font.name = 'Times New Roman'
                else:
                    if r_i % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    for p_cell in cell.paragraphs:
                        for r_cell in p_cell.runs:
                            r_cell.font.size = Pt(9.5)
                            r_cell.font.name = 'Times New Roman'
        doc.add_paragraph() # space after table

    def flush_code(c_lines):
        if not c_lines: return
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(6)
        p_c.paragraph_format.space_after = Pt(6)
        p_c.paragraph_format.line_spacing = 1.0
        r_c = p_c.add_run("\n".join(c_lines))
        r_c.font.name = 'Consolas'
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = RGBColor(30, 41, 59)

    for l in lines[start_idx:]:
        line = l.rstrip()

        # Handle page breaks
        if "\\pagebreak" in line or line.strip() == "---":
            if in_table:
                flush_table(table_lines)
                in_table = False
                table_lines = []
            if in_code:
                flush_code(code_lines)
                in_code = False
                code_lines = []
            if "\\pagebreak" in line:
                doc.add_page_break()
            continue

        # Handle Code blocks
        if line.strip().startswith("```"):
            if in_code:
                flush_code(code_lines)
                in_code = False
                code_lines = []
            else:
                if in_table:
                    flush_table(table_lines)
                    in_table = False
                    table_lines = []
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        # Handle Tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            continue
        elif in_table:
            flush_table(table_lines)
            in_table = False
            table_lines = []

        # Handle Headings
        if line.startswith("# "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)
            r = h.add_run(line[2:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor(160, 20, 30)
        elif line.startswith("## "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(line[3:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138) # Deep Blue
        elif line.startswith("### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            r = h.add_run(line[4:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(55, 65, 81)
        elif line.startswith("#### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(line[5:].strip())
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.italic = True
        elif line.strip().startswith("- ") or line.strip().startswith("* "):
            p_bullet = doc.add_paragraph(style='List Bullet')
            p_bullet.paragraph_format.space_after = Pt(3)
            # Simple bold parsing
            text = line.strip()[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p_bullet.add_run(part[2:-2]).bold = True
                else:
                    p_bullet.add_run(part)
        elif line.strip() == "":
            pass
        else:
            p_text = doc.add_paragraph()
            p_text.paragraph_format.line_spacing = 1.15
            p_text.paragraph_format.space_after = Pt(6)
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p_text.add_run(part[2:-2]).bold = True
                else:
                    p_text.add_run(part)

    # Save to .docx
    output_path = "docs/AttendX_PBL_Project_Report.docx"
    doc.save(output_path)
    print(f"Successfully generated Word Document at: {output_path}")

if __name__ == "__main__":
    create_report()
